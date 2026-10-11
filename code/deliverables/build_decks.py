"""Build the internal and client decks for the Revantage Luxembourg workshop.

Uses the vendored databricks-slides generator (slide_kit/, Databricks template) and adds three
slide types it lacks: `screenshot` (title + live screenshot, optional bullets), `demo` (internal
per-demo slide: problem / build / value / why Databricks + screenshot) and `build-step` (client
build-along slide: the prompt you type, the checkpoint, the fallback + screenshot).

Usage (from the project root):
    uv run --with python-pptx python code/deliverables/build_decks.py
"""
import sys
from pathlib import Path

from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE / "slide_kit" / "scripts"))
import generate_pptx as g  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402

SHOTS = ROOT / "notes" / "agent-notes" / "screenshots"
ASSETS = ROOT / "assets"
NAME = "client__revantage__luxembourg-workshop__2026-10-08__tbd"

for t in ("screenshot", "demo", "build-step"):
    g.VALID_SLIDE_TYPES.add(t)
    g.LAYOUT_MAPPINGS[t] = ["CUSTOM"]
C = g.COLORS
DARK, MUTED, ACCENT = C["text_dark"], "#5A6B73", C["accent"]


def shot(name):
    p = ASSETS / name if (ASSETS / name).exists() else SHOTS / name
    if not p.exists():
        raise FileNotFoundError(p)
    return p


class Deck(g.DatabricksSlideGenerator):
    def _title(self, slide, text, subtitle=None):
        self.add_textbox(slide, text, 0.6, 0.35, 12.1, 0.7, font_size=28, bold=True, color=DARK)
        if subtitle:
            self.add_textbox(slide, subtitle, 0.6, 0.98, 12.1, 0.45, font_size=14, color=MUTED)

    def _image(self, slide, path, left, top, width, height):
        """Fit an image inside the box, keeping its aspect ratio, with a thin border."""
        w, h = Image.open(path).size
        scale = min(width / w, height / h)
        iw, ih = w * scale, h * scale
        pic = slide.shapes.add_picture(str(path), Inches(left + (width - iw) / 2), Inches(top),
                                       Inches(iw), Inches(ih))
        pic.line.color.rgb = g.hex_to_rgb("#D0D5D8")
        pic.line.width = Pt(0.75)

    def _block(self, slide, label, text, left, top, width, height, size=12, label_color=ACCENT):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = label.upper()
        p.font.size, p.font.bold, p.font.name = Pt(10), True, g.FONT_FAMILY
        p.font.color.rgb = g.hex_to_rgb(label_color)
        for line in (text if isinstance(text, list) else [text]):
            q = tf.add_paragraph()
            q.text = line
            q.font.size, q.font.name = Pt(size), g.FONT_FAMILY
            q.font.color.rgb = g.hex_to_rgb(DARK)
            q.space_before = Pt(3)

    def add_screenshot_slide(self, d):
        s = self._create_slide("screenshot", d)
        self._title(s, d["title"], d.get("subtitle"))
        top = 1.5 if d.get("subtitle") else 1.2
        if d.get("bullets"):
            self._block(s, d.get("label", "What you see"), d["bullets"], 0.6, top, 3.9, 5.5, size=13)
            self._image(s, shot(d["image"]), 4.7, top, 8.1, 5.6)
        else:
            self._image(s, shot(d["image"]), 0.6, top, 12.1, 5.75)
        if d.get("caption"):
            self.add_textbox(s, d["caption"], 0.6, 6.9, 12.1, 0.35, font_size=11, color=MUTED)

    def add_demo_slide(self, d):
        s = self._create_slide("demo", d)
        self._title(s, d["title"], d.get("subtitle"))
        w = 4.3 if d.get("image") else 5.9
        y = 1.5
        for label, key in (("Problem", "problem"), ("They build", "build"),
                           ("Value", "value"), ("Why Databricks", "why")):
            if d.get(key):
                n = 1 + len(d[key]) // 62
                self._block(s, label, d[key], 0.6, y, w, 0.5 + 0.22 * n)
                y += 0.62 + 0.22 * n
        if d.get("watch"):
            self._block(s, "Facilitator watch-out", d["watch"], 0.6, y, w, 1.0, size=11,
                        label_color="#B45309")
        if d.get("image"):
            self._image(s, shot(d["image"]), 5.1, 1.5, 7.7, 5.5)
        elif d.get("right"):
            self._block(s, d.get("right_label", "Key numbers"), d["right"], 6.8, 1.5, 5.9, 5.0, size=13)

    def add_build_step_slide(self, d):
        s = self._create_slide("build-step", d)
        self._title(s, d["title"], d.get("subtitle"))
        w = 4.6 if d.get("image") else 12.1
        self.add_textbox(s, "YOU TYPE (Genie Code)" if not d.get("prompt_label") else d["prompt_label"],
                         0.6, 1.5, w, 0.3, font_size=10, bold=True, color=ACCENT)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.85),
                                 Inches(w), Inches(1.9))
        box.fill.solid()
        box.fill.fore_color.rgb = g.hex_to_rgb("#EEF2F3")
        box.line.color.rgb = g.hex_to_rgb("#C9D2D6")
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.15)
        p = tf.paragraphs[0]
        p.text = d["prompt"]
        p.alignment = PP_ALIGN.LEFT
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p.font.size, p.font.name = Pt(13), "Courier New"
        p.font.color.rgb = g.hex_to_rgb(DARK)
        self._block(s, "Checkpoint", d["checkpoint"], 0.6, 3.95, w, 1.4)
        if d.get("stuck"):
            self._block(s, "If you're stuck", d["stuck"], 0.6, 5.45, w, 1.2, size=11, label_color=MUTED)
        if d.get("image"):
            self._image(s, shot(d["image"]), 5.4, 1.5, 7.4, 5.5)

    def generate(self, content):
        extra = {"screenshot": self.add_screenshot_slide, "demo": self.add_demo_slide,
                 "build-step": self.add_build_step_slide}
        for sd in content["slides"]:
            if sd.get("type") in extra:
                extra[sd["type"]](sd)
            else:
                g.DatabricksSlideGenerator.generate(self, {"slides": [sd]})


AGENDA = ["Day 1 AM · Genie One, Genie Code, document agents",
          "Day 1 PM · Metric views, dashboards, Apps, governance",
          "Day 2 AM · Lakeflow, Lakebase, AI Gateway, forecasting",
          "Day 2 PM · Build-your-own lab, Genie accuracy, domains & Pages, data quality"]

STORY = ["Southeast net profit fell 18.6% in 2026-Q2 (€2.58M → €2.09M)",
         "Almost all of it is Harbor Tower: anchor tenant left on 31 March",
         "…and a USD 339,600 chiller (€312,390.76) booked to OpEx; the policy says CapEx",
         "Luxembourg City grew fastest (+10.0% QoQ net profit)",
         "Every demo uncovers one more piece of the same story"]

# ---------------------------------------------------------------------------------------------
# Internal enablement deck (SA / SSA / AE who will re-deliver the workshop)
# ---------------------------------------------------------------------------------------------
INTERNAL = {"title": "Revantage Luxembourg Workshop", "slides": [
    {"type": "title", "title": "Revantage build-along workshop",
     "subtitle": "Internal enablement: how to re-deliver the 18 demos", "date": "October 2026"},
    {"type": "stat-row", "title": "Verified before day one", "stats": [
        {"value": "18", "label": "demos on one dataset"},
        {"value": "18", "label": "deployed answer keys (ref_*)"},
        {"value": "100%", "label": "participant prompts run in the UI"}]},
    {"type": "agenda", "title": "Two days, four blocks", "items": AGENDA},
    {"type": "three-column", "title": "The build-live model",
     "columns": [
         {"header": "Starting state", "items": ["Bundle job loads data, PDFs, the broken Snowflake SQL",
                                                "Each participant gets <my_schema>", "Reset scripts per demo"]},
         {"header": "Answer key (ref_*)", "items": ["Finished asset deployed by the bundle",
                                                    "Fallback when a live build stalls",
                                                    "Source of truth for checkpoint numbers"]},
         {"header": "Build-along guide", "items": ["Exact prompt per stage, time-boxed",
                                                   "Expected result + checkpoint",
                                                   "\"If stuck\" path + facilitator watch-outs"]}]},
    {"type": "content", "title": "One dataset tells one story", "bullets": STORY,
     "notes": "Planted facts F1–F7 are asserted by the generator (code/data-gen/generate.py) and re-checked by tests/check_starting_state.py (28 checks)."},
    {"type": "screenshot", "title": "Architecture: every asset the bundle deploys",
     "image": "architecture.png", "caption": "Source: assets/architecture.drawio (code/deliverables/architecture.py)"},
    {"type": "section", "title": "Day 1", "subtitle": "Ask, build, extract, define, present, govern"},
    {"type": "demo", "title": "D2 · Genie One: ask the portfolio", "image": "d02-genie-southeast.png",
     "problem": "Finance answers \"why did X move?\" with spreadsheets and analyst queues.",
     "build": "Participants teach the Start agent (instructions, sample questions) until it answers like the Reference.",
     "value": "Root cause in one question: Southeast −€480K, all Harbor Tower (−44%).",
     "why": "Genie reasons over UC-governed tables and metric views; answers show their SQL and respect permissions.",
     "watch": "Set the default-period instruction (2026-Q2, QoQ = Q2 vs Q1) or Genie picks the wrong quarter."},
    {"type": "demo", "title": "D3 · Genie Code: migrate a Snowflake view", "image": "s9-d03-genie-code.png",
     "problem": "Migration fear: hundreds of Snowflake views, unknown breakage.",
     "build": "Paste the Snowflake view; Genie Code fixes NUMBER, TO_VARCHAR, DIV0, SYSDATE, then finds the region fan-out join (176 → 64 rows) and adds YoY.",
     "value": "Most Snowflake SQL runs unchanged; the assistant finds the real bugs, including one that was already wrong in Snowflake.",
     "why": "Genie Code works in the editor with catalog context and runs/validates what it writes.",
     "watch": "The inline Quick Fix proposes CAST AS NUMBER (wrong). Reject it and use the chat."},
    {"type": "demo", "title": "D4–D5 · Agent Bricks: invoices and policies", "image": "d04-extraction.png",
     "problem": "Invoices are PDFs; policy answers live in people's heads.",
     "build": "Information Extraction agent over 8 invoices (72/72 fields), reconcile to GL; Knowledge Assistant over 4 policies.",
     "value": "The USD 339,600 chiller is found booked to OpEx; the policy says CapEx and needs CFO + AM Director.",
     "why": "Agents are built from UI, evaluated, and called from SQL (ai_query) next to governed data.",
     "watch": "KA got 6/8 approver tiers right in the batch prompt; teach verify-against-a-rule (D5 Stage 4)."},
    {"type": "demo", "title": "D7 · Metric views: define net profit margin once", "image": "d07-metric-view.png",
     "problem": "Finance (59.3%) and Asset Management (54.0%) report different margins.",
     "build": "Three UC metric views (margin, target attainment, QoQ growth) with YAML measures.",
     "value": "One approved definition used by dashboards, Genie and SQL alike.",
     "why": "Metric views are catalog objects: governed, certified, lineage-tracked, queried with MEASURE().",
     "watch": "Stage 3 checkpoint is net-profit growth (+10.0% Lux City); revenue growth is +9.2%."},
    {"type": "demo", "title": "D8 · AI/BI dashboard: the quarterly review pack", "image": "d08-p1-executive.png",
     "problem": "The board pack is Excel and emailed PDFs.",
     "build": "4-page dashboard built with Genie Code in the dashboard editor, global filters, Ask Genie.",
     "value": "€10.7M · 59.3% · 62.5% live, with Genie for the follow-up question.",
     "why": "Dashboards sit on metric views, publish with embedded credentials, and include Genie.",
     "watch": "Prompts must say \"run it and check\": Genie Code once reported a failing dataset as done."},
    {"type": "demo", "title": "D9 · Apps: lease expiry tracker", "image": "d09-app.png",
     "problem": "Lease roll is tracked in a spreadsheet nobody trusts.",
     "build": "From a Streamlit template, Genie Code adds the query, filters, risk column, cards and CSV export.",
     "value": "6 / 14 / 23 leases in 90 / 180 / 365 days; €103k monthly rent at risk.",
     "why": "Apps run on serverless next to the data, authenticate with Databricks identity, no per-seat licences.",
     "watch": "Name the warehouse resource (sql-warehouse) in the prompt or app.yaml wiring breaks."},
    {"type": "demo", "title": "D10 · Unity Catalog governance", "image": "d10-masked-pii.png",
     "problem": "Who can see tenant emails and phone numbers, and how do we prove it?",
     "build": "Business descriptions, certified tag, column masks with is_account_group_member, PII class tags.",
     "value": "Same table, different views by group; certification visible in search and Genie.",
     "why": "One governance layer for tables, metric views, models, agents and apps.",
     "watch": "Prompt must name class.email_address / class.phone_number; Genie Code tried to edit the account tag policy.",
     "right": ["contact_email → l***@nimbus.example.com", "contact_phone → ***597",
               "system.certification_status = certified", "Reset script returns the starting state"]},
    {"type": "section", "title": "Day 2", "subtitle": "Engineer, operate, monitor, forecast, trust"},
    {"type": "demo", "title": "D11 · Lakeflow: bronze → silver → gold", "image": "d11-pipeline-graph.png",
     "problem": "Nightly CSV extracts, silent bad data, no lineage.",
     "build": "Genie Code (pipeline editor) designs 3 SQL files, adds expectations; a late SAP-LUX batch with \"1.850,75\" fails the update; Genie Code diagnoses and fixes it.",
     "value": "Gold matches the oracle on all 80 rows; bad money never reaches the board pack.",
     "why": "Declarative pipelines with expectations, lineage, retries and Auto Loader in one object.",
     "watch": "Prompt must say <entity>*.csv and amount as STRING, or the late batch is silently ignored."},
    {"type": "demo", "title": "D12 · Lakebase: approval workflow", "image": "d12-approvals-app.png",
     "problem": "Two-signature approvals happen in email; the GL can't tell approved from pending.",
     "build": "Postgres schema with CHECK + generated column; the app enforces the second approver; a job MERGEs approved rows to silver.",
     "value": "Only approved invoices reach the lakehouse; every action audited.",
     "why": "Managed Postgres next to Delta: OAuth app identity, synced tables for reverse ETL.",
     "watch": "Lakebase SQL lives in the SQL editor (compute → More… → Lakebase Postgres)."},
    {"type": "demo", "title": "D13 · AI Gateway: usage and cost", "image": "d13-ai-gateway.png",
     "problem": "Nobody knows who uses which model and what it costs.",
     "build": "Usage, error, cost queries on system tables and an hourly alert on daily tokens.",
     "value": "Vector Search for an idle KA index was the top AI cost (~$14 of ~$19).",
     "why": "Usage and payloads are UC tables; rate limits and guardrails live on the gateway.",
     "watch": "Pay-per-token endpoints report to system.serving.endpoint_usage, not system.ai_gateway.usage."},
    {"type": "demo", "title": "D14 · Forecast with ai_forecast", "image": "d08-p4-forecast.png",
     "problem": "Quarterly reforecast takes weeks of spreadsheet work.",
     "build": "One SQL statement forecasts every property to year end and joins targets.",
     "value": "Harbor Tower and Zuidas flagged below 2026-Q3 target.",
     "why": "Forecasting as a governed SQL function on a warehouse; no model objects to manage.",
     "watch": "Genie Code says ai_forecast is \"disabled\" (its probe runs on serverless). Run on the warehouse."},
    {"type": "demo", "title": "D16 · Data quality: trust scorecard", "image": "d16-outlier-flagged.png",
     "problem": "A 10× revenue load would reach every board number before anyone notices.",
     "build": "Profiling monitor, leave-one-out outlier rule, lineage blast radius, alert.",
     "value": "0 outliers on clean data, exactly 1 (z ≈ 172) after the seed.",
     "why": "Monitoring, lineage and alerts on the same catalog as the data.",
     "watch": "Ask for leave-one-out σ explicitly; inclusive σ caps z near 3 with 11 quarters.",
     "right": ["Seed: Harbor Tower 2026-Q3 revenue €17.2M, NULL net_profit",
               "Outlier rule: z > 3 vs other quarters AND > 50% from mean",
               "Blast radius: metric views, dashboard, Genie agents, forecast",
               "Reset: code/sql/d16_reset.sql"]},
    {"type": "demo", "title": "D17 · Genie accuracy ladder", "image": "d17-evaluations-history.png",
     "problem": "A Genie Agent on the legacy warehouse gets 2 of 10 business questions right.",
     "build": "Clone the Start agent, then add names, comments, descriptions + keys, metric views, instructions; re-run the benchmark after each.",
     "value": "Genie's own benchmark: 20% → 10% → 30% → 50% → 100%; every failure explained, with Attempt fix.",
     "why": "Accuracy is engineered in Unity Catalog metadata that every tool reuses, and measured, not felt.",
     "watch": "Set Mode: Chat before each run (resets to Agent; 3 vs 8 min). CREATE OR REPLACE keeps old comments.",
     "right": ["Names alone: 10% (codes R/O/M/C/I still undecoded)",
               "Comments: 30% (codes decoded, rules guessed)",
               "Descriptions + PK/FK: 50% (lease history understood)",
               "Metric views: 100% (approved net profit)"]},
    {"type": "demo", "title": "D18 · Domains, Pages & Metrics", "image": "d18-page-genie-code.png",
     "problem": "Genie One sees every \"net profit\" in the account and has to guess which one you mean.",
     "build": "A Page (definition, synonyms, related metric view) written by Genie Code in the Revantage domain.",
     "value": "Genie One reads the Page (\"NOI … per the glossary page\") and names competing sources.",
     "why": "Domains organize, Pages explain, metric views calculate, certification says what to trust.",
     "watch": "Pages feed Genie One, not Genie Agents (agent 50% with or without). Conflicting certified views win: fix trust, not prompts."},
    {"type": "two-column", "title": "What we learned running every prompt live",
     "left_header": "Genie Code patterns", "left": [
         "Ends every build prompt with \"run it and check\"",
         "Names exact tables, tags and resource keys",
         "Spells out the method (leave-one-out, units)",
         "Pauses for approval on destructive DDL: show it"],
     "right_header": "Facilitator rules", "right": [
         "Run pre-flight in the Revantage workspace",
         "Use reset scripts between groups",
         "Fall back to ref_* the moment a build stalls",
         "Never let an agent change account-level policy"]},
    {"type": "checklist", "title": "Before the workshop", "items": [
        {"text": "Revantage workspace access + bundle target filled (revantage)", "checked": False},
        {"text": "Pre-flight: preflight.py, system table grants, D16 monitor, D2/D9 show-items", "checked": False},
        {"text": "Setup job + participant sandboxes run; starting-state check ALL PASS", "checked": False},
        {"text": "All 18 references deployed and verified on FEVM (D18 domain: Revantage workspace)", "checked": True},
        {"text": "D18: run code/scripts/d18_domain.sh, publish domain + reference Pages, grant MANAGE", "checked": False},
        {"text": "Every participant prompt run in the Genie Code UI", "checked": True}]},
    {"type": "closing", "title": "Thank you"},
]}

# ---------------------------------------------------------------------------------------------
# Client deck (Revantage participants): build-along structure, no competitive positioning
# ---------------------------------------------------------------------------------------------
CLIENT = {"title": "Revantage Luxembourg Workshop", "slides": [
    {"type": "title", "title": "Build it yourself",
     "subtitle": "Two-day hands-on workshop · Luxembourg", "date": "October 2026"},
    {"type": "section", "title": "You build it. Genie Code guides."},
    {"type": "agenda", "title": "Two days, four blocks", "items": AGENDA},
    {"type": "timeline", "title": "How each session works", "steps": [
        {"title": "Goal", "description": "One business question, one asset"},
        {"title": "Prompt", "description": "Type the prompt into Genie Code"},
        {"title": "Checkpoint", "description": "Compare with the expected number"},
        {"title": "If stuck", "description": "Open the finished reference and continue"}]},
    {"type": "content", "title": "Your data, one story", "bullets": STORY},
    {"type": "screenshot", "title": "What we'll build together", "image": "architecture.png"},
    {"type": "section", "title": "Day 1", "subtitle": "Ask, build, extract, define, present, govern"},
    {"type": "build-step", "title": "Ask your portfolio a question (Genie One)", "image": "d02-genie-southeast.png",
     "prompt": "Why did Southeast net profit drop in 2026-Q2? Break it down by property.",
     "checkpoint": "Southeast −€480K (−19%); Harbor Tower −€507K (−44%); the other two properties up.",
     "stuck": "Ask the same question in Revantage Portfolio Q&A (Reference)."},
    {"type": "build-step", "title": "Migrate a Snowflake view (Genie Code)", "image": "s9-d03-genie-code.png",
     "prompt": "Explain this query and identify any issues for Databricks SQL. It was written for Snowflake.",
     "checkpoint": "4 functions to change; then a join bug: 176 rows → 64 after the fix. Harbor Tower YoY −478,309.66.",
     "stuck": "Compare with ref_d03_property_performance_v."},
    {"type": "build-step", "title": "Turn invoices into a table (Agent Bricks)", "image": "d04-extraction.png",
     "prompt_label": "YOU DO (Agent Bricks UI) then ask Genie Code",
     "prompt": "Explain this query step by step, then wrap it in CREATE OR REPLACE TABLE <my_schema>.invoices_extracted.",
     "checkpoint": "8 invoices; HV-2026-0417 USD 339,600 matches a €312,390.76 GL posting in Repairs & Maintenance.",
     "stuck": "Run code/reference/d04_extract_invoices.sql."},
    {"type": "build-step", "title": "One definition of net profit margin", "image": "d07-metric-view.png",
     "prompt": "Create a Unity Catalog metric view <my_schema>.mv_net_profit_margin over gold_property_quarterly … plus asset_mgmt_net_profit_margin labelled as not approved.",
     "checkpoint": "2026-Q2: 59.33% approved vs 53.95% Asset Management definition.",
     "stuck": "Copy the first statement of code/reference/d07_metric_views.sql."},
    {"type": "build-step", "title": "The quarterly review pack (AI/BI)", "image": "d08-p1-executive.png",
     "prompt": "On a page \"Executive Summary\", add counters for Revenue 2026-Q2 (EUR, compact), Net profit margin (%), Target attainment (%) …",
     "checkpoint": "€10.7M · 59.3% · 62.5%; the line chart shows Southeast dropping in Q2.",
     "stuck": "Open Revantage Portfolio Review (Reference) → Data tab."},
    {"type": "build-step", "title": "A lease expiry tracker (Apps)", "image": "d09-app.png",
     "prompt": "Change this app to query silver_leases joined to silver_tenants and silver_properties … for leases ending in the next 90 days.",
     "checkpoint": "6 leases in 90 days; €103k monthly rent at risk; 14 in 180 days.",
     "stuck": "Open revantage-lease-tracker and compare its query."},
    {"type": "build-step", "title": "Who can see what (Unity Catalog)", "image": "d10-masked-pii.png",
     "prompt": "Create a column mask function that shows the full email only to members of revantage_finance_stewards … Tag the columns with class.email_address and class.phone_number.",
     "checkpoint": "You see l***@nimbus.example.com and ***597; stewards see full values.",
     "stuck": "Run code/reference/d10_governance.sql in your schema."},
    {"type": "section", "title": "Day 2", "subtitle": "Engineer, operate, monitor, forecast, trust"},
    {"type": "build-step", "title": "Bronze → silver → gold (Lakeflow)", "image": "d11-pipeline-graph.png",
     "prompt": "Design a Bronze → Silver → Gold Lakeflow Declarative Pipeline in SQL … match <entity>*.csv … read gl_transactions.amount as STRING …",
     "checkpoint": "Gold 80 rows equal to the certified table; a late batch with \"1.850,75\" stops the pipeline.",
     "stuck": "Copy code/pipelines/ref_property_pipeline/."},
    {"type": "build-step", "title": "Two-signature approvals (Lakebase)", "image": "d12-approvals-app.png",
     "prompt": "Create a Postgres schema approvals with a table invoice_approvals … requires_two_approvers = total_eur > 50000 … The final approver must differ from the first approver.",
     "checkpoint": "€312k invoice needs two people; the same person can't approve twice; only approved invoices reach silver.",
     "stuck": "code/lakebase/d12_invoice_approvals.sql."},
    {"type": "build-step", "title": "Who's using AI, and what it costs", "image": "d13-ai-gateway.png",
     "prompt": "Using system.ai_gateway.usage, show total requests, total tokens and average latency by endpoint for workspace :workspace_id over the last 30 days.",
     "checkpoint": "Every AI call by model, person and cost; an alert before daily tokens exceed 1M.",
     "stuck": "code/reference/d13_ai_gateway_queries.sql."},
    {"type": "build-step", "title": "Where is revenue heading? (ai_forecast)", "image": "d08-p4-forecast.png",
     "prompt": "Forecast quarterly revenue for all properties through 2026-12-31 with ai_forecast … join to silver_quarterly_targets … below revenue_target.",
     "checkpoint": "Harbor Tower and Zuidas Residences forecast below their 2026-Q3 targets.",
     "stuck": "ref_d14_revenue_forecast."},
    {"type": "build-step", "title": "Trust scorecard (data quality)", "image": "d16-outlier-flagged.png",
     "prompt": "Flag property-quarters whose revenue is more than 3 standard deviations from the mean of that property's OTHER quarters (leave-one-out) and more than 50% away from that mean.",
     "checkpoint": "0 rows on clean data; exactly 1 after the bad load (Harbor Tower 2026-Q3).",
     "stuck": "code/reference/d16_revenue_outliers.sql."},
    {"type": "build-step", "title": "From 2 to 10 out of 10 (Genie accuracy)", "image": "d17-evaluations-history.png",
     "prompt": "Using legacy_data_dictionary, create one table in <my_schema> for each legacy_* table … rename every column to its business_name. Then add the column comments, table descriptions and keys it describes.",
     "checkpoint": "Benchmark after each layer: 20% → 10% → 30% → 50% → 100% with metric views and instructions.",
     "stuck": "Revantage Accuracy Lab (Reference)."},
    {"type": "build-step", "title": "Teach Genie One your words (Pages)", "image": "d18-page-genie-code.png",
     "prompt": "Help me create a page \"Net profit\" in the Revantage Workshop domain … CapEx and intercompany are not deducted … \"This quarter\" means 2026-Q2 … Related assets: ref_mv_net_profit_margin.",
     "checkpoint": "Genie One reads the Page: NOI = net profit, this quarter = 2026-Q2.",
     "stuck": "code/reference/d18_pages.md."},
    {"type": "stat-row", "title": "What you'll leave with", "stats": [
        {"value": "18", "label": "assets you built yourself"},
        {"value": "1", "label": "governed dataset behind all of them"},
        {"value": "2", "label": "days, hands on keyboard"}]},
    {"type": "three-column", "title": "Outcomes for Revantage",
     "columns": [
         {"header": "Finance & FP&A", "items": ["Answers in minutes, not analyst queues",
                                                "One approved definition of every KPI",
                                                "Forecast every property in one query"]},
         {"header": "Asset Management", "items": ["Lease roll and rent at risk, always current",
                                                  "Invoice and policy checks automated",
                                                  "Two-signature approvals enforced"]},
         {"header": "Data & IT", "items": ["Migration path from existing SQL",
                                           "Quality rules that stop bad loads",
                                           "AI usage and cost visible and governed"]}]},
    {"type": "timeline", "title": "Proposed next steps (to agree together)", "steps": [
        {"title": "Choose", "description": "Pick two workshop assets to take further"},
        {"title": "Connect", "description": "Point them at real Yardi / SAP extracts"},
        {"title": "Pilot", "description": "Share with a Finance pilot group"},
        {"title": "Review", "description": "Decide on the next use cases"}]},
    {"type": "closing", "title": "Thank you"},
]}


def build(content, out):
    d = Deck(g.TEMPLATE_PATH)
    d.generate(content)
    d.save(str(out))
    print("wrote", out, len(d.prs.slides.__iter__.__self__._sldIdLst), "slides")


if __name__ == "__main__":
    build(INTERNAL, ASSETS / f"{NAME}-internal.pptx")
    build(CLIENT, ASSETS / f"{NAME}-client.pptx")

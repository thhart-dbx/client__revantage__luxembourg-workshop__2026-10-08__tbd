"""Build the module-by-module build-along deck (one module per demo).

Per demo:  section slide → "What we'll build" (overview + finished asset) → "How we'll get there"
(the steps) → one slide per step / per prompt: the prompt to type (or the click-path instruction for
steps without one), the expected result, the fallback, and a screenshot of the expected output.

Text comes from the verified guides in code/prompts/ (facilitator-only notes are dropped); the
screenshots come from code/deliverables/module_manifest.py. Usage (from the project root):
    uv run --with python-pptx --with pillow --with pygments python code/deliverables/build_modules_deck.py
"""
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
import build_decks as bd  # noqa: E402  (Deck class, template, slide kit)
from build_docs import FACILITATOR, parse  # noqa: E402
from module_manifest import OVERVIEW, STEPS, fname  # noqa: E402

g = bd.g
from pptx.util import Inches, Pt  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402
from pygments.lexers import get_lexer_by_name  # noqa: E402
from pygments.token import Comment, Keyword, Name, Number, Operator, String  # noqa: E402

SHOTS = ROOT / "notes" / "agent-notes" / "screenshots"
MONO = "Courier New"
CODE_INK = "#0B4F6C"  # inline `code` in running text
# SQL syntax colours (token type → colour, bold, italic)
HL = [(Comment, "#7A869A", False, True), (String, "#2E7D32", False, False), (Number, "#B5651D", False, False),
      (Keyword, "#0B5CAD", True, False), (Name.Builtin, "#8E3BA8", False, False),
      (Operator, "#5A6B73", False, False)]
CODE_FILE = re.compile(r"`?(code/[\w./-]+\.(?:sql|py|sh))`?")
# one line per module for the "Workshop at a glance" slide
TAGLINE = {1: "Catalog, SQL editor, Genie Code: your bearings", 2: "Ask the portfolio questions in plain English",
           3: "Migrate and fix a broken Snowflake view", 4: "Invoice PDFs to a Delta table",
           5: "Cited answers from the group policies", 6: "Audit the data before you trust it",
           7: "Define Net Profit Margin once", 8: "Executive Portfolio Review dashboard",
           9: "Lease Expiry Tracker app", 10: "Business metadata, certification, PII masks",
           11: "Raw CSVs to bronze, silver, gold", 12: "Invoice approvals on Lakebase, approved to Delta",
           13: "AI usage and cost queries, plus an alert", 14: "Revenue forecast for Q3/Q4 vs target",
           15: "Pick a track, build your own", 16: "Monitoring, outlier alert, blast radius",
           17: "Genie accuracy, one layer at a time", 18: "Business definitions Genie One can use"}
for t in ("module-step", "module-overview", "module-steps", "module-agenda"):
    g.VALID_SLIDE_TYPES.add(t)
    g.LAYOUT_MAPPINGS[t] = ["CUSTOM"]


def find_image(name):
    for d in (SHOTS / "modules", SHOTS, ROOT / "assets"):
        if (d / name).exists():
            return d / name
    raise FileNotFoundError(name)


def clean(s, code=False):
    """Strip markdown. code=True keeps `inline code` spans (rendered monospaced by rich())."""
    s = re.sub(r"\s*\([^)]*FEVM[^)]*\)", "", s)
    s = re.split(r"\*?\*?FEVM note|On the FEVM", s)[0]  # workspace-specific numbers stay out of the participant deck
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s).replace("**", "")
    s = s.replace("\\", "") if code else s.replace("`", "").replace("\\", "")
    s = re.sub(r"\[(.+?)\]\([^)]*\)", r"\1", s)
    return re.sub(r"\s+", " ", s).strip()


def clip(s, n):
    return s if len(s) <= n else s[: n - 1].rsplit(" ", 1)[0] + " …"


def stage_key(title):
    m = re.match(r"(Stage \d+|Track [A-Z]|Closing)", title)
    return m.group(1) if m else title.split(":")[0]


def short(title):
    t = re.sub(r"^(Stage \d+( \([^)]*\))?|Track [A-Z])\s*[:—-]\s*", "", title)
    return re.sub(r"\s*\((\d+ min[^)]*|facilitator[^)]*|optional[^)]*)\)\s*$", "", t).strip()


def steps_of(stage, guide):
    """Split a stage into steps: one per prompt, or one 'do' step if there is no prompt.

    Code blocks keep their line breaks (rendered highlighted); `source` is the file that holds the
    full code: a code/… file named in the stage, else the guide itself."""
    lines = [l for l in stage["lines"] if not FACILITATOR.search(l)]
    prompts, cur, intro, stuck, code, in_code, want_code = [], None, [], [], [], False, None
    named = CODE_FILE.search("\n".join(lines))
    source = named.group(1) if named else f"code/prompts/{guide.name}"
    for l in lines:
        if l.startswith("```"):
            in_code = not in_code
            if not in_code and want_code is not None:
                want_code.append("`" + " ".join(c.strip() for c in code[:3]) + "`")
                want_code, code = None, []
            elif not in_code:
                if cur is None:
                    intro.append(("CODE", "\n".join(code).strip("\n")))
                code = []
            continue
        if in_code:
            code.append(l)
            continue
        if l.startswith(">"):
            if cur is None or cur["expected"]:
                cur = {"prompt": clean(l.lstrip("> ")), "expected": []}
                prompts.append(cur)
            else:
                cur["prompt"] += " " + clean(l.lstrip("> "))
        elif re.match(r"\*\*(Expected|Checkpoint)", l):
            body = clean(re.sub(r"^\*\*[^*]+\*\*:?\s*", "", l), code=True)
            if cur is None:
                intro.append(("EXP", body))
            elif body:
                cur["expected"].append(body)
            if not body or body.endswith(":"):
                want_code = cur["expected"] if cur else None
        elif l.startswith("**If stuck"):
            stuck.append(clean(re.sub(r"^\*\*[^*]+\*\*:?\s*", "", l), code=True))
        elif l.startswith("**Prompt"):
            continue
        elif l.strip() and not l.startswith("|") and cur is None:
            intro.append(("TXT", clean(re.sub(r"^\s*(?:[-*]|\d+\.)\s+", "", l), code=True)))
    if not prompts:
        txt = [t for k, t in intro if k == "TXT" and t]
        sql = [t for k, t in intro if k == "CODE" and t]
        exp = [t for k, t in intro if k == "EXP" and t]
        return [{"prompt": None, "do": clip(" ".join(txt[:3]), 420), "sql": sql[0] if sql else None,
                 "source": source, "expected": " ".join(exp), "stuck": " ".join(stuck)}]
    return [{"prompt": p["prompt"], "expected": " ".join(p["expected"]), "stuck": " ".join(stuck),
             "do": None, "sql": None, "source": source} for p in prompts]


def rich(p, text, size, color=None):
    """Fill a paragraph: plain text as is, `inline code` monospaced in CODE_INK."""
    if text.count("`") % 2:  # a clipped span: drop the dangling backtick
        text = "".join(text.rsplit("`", 1))
    p.alignment = PP_ALIGN.LEFT
    for i, part in enumerate(re.split(r"`([^`]*)`", text)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.name = MONO if i % 2 else g.FONT_FAMILY
        r.font.color.rgb = g.hex_to_rgb(CODE_INK if i % 2 else (color or bd.DARK))


def highlight(p, line, size):
    """One line of SQL as coloured runs; <placeholders> stay one plain token."""
    p.alignment = PP_ALIGN.LEFT
    for k, seg in enumerate(re.split(r"(<\w+>|\$\{\w+\}|(?<![:\w]):\w+)", line)):
        if k % 2:
            r = p.add_run()
            r.text = seg
            r.font.size, r.font.name, r.font.italic = Pt(size), MONO, True
            r.font.color.rgb = g.hex_to_rgb(bd.ACCENT)
        elif seg:
            _tokens(p, seg, size)


def _tokens(p, line, size):
    for ttype, val in get_lexer_by_name("sql").get_tokens(line):
        val = val.rstrip("\n")
        if not val:
            continue
        r = p.add_run()
        r.text = val
        r.font.size, r.font.name = Pt(size), MONO
        colour, bold, italic = next(((c, b, i) for t, c, b, i in HL if ttype in t), (bd.DARK, False, False))
        r.font.color.rgb = g.hex_to_rgb(colour)
        r.font.bold, r.font.italic = bold, italic


class ModuleDeck(bd.Deck):
    def add_module_overview_slide(self, d):
        s = self._create_slide("module-overview", d)
        self._title(s, d["title"], d.get("subtitle"))
        y = 1.55
        for label, text in d["blocks"]:
            n = 1 + len(text) // 58
            self._block(s, label, text, 0.6, y, 4.3, 0.45 + 0.22 * n, size=12)
            y += 0.6 + 0.22 * n
        self._image(s, find_image(d["image"]), 5.1, 1.5, 7.7, 5.5)

    def add_module_steps_slide(self, d):
        s = self._create_slide("module-steps", d)
        self._title(s, d["title"], d.get("subtitle"))
        items = d["items"]
        per_col = (len(items) + 1) // 2 if len(items) > 5 else len(items)
        for i, (name, kind) in enumerate(items):
            col, row = divmod(i, per_col)
            x, y = 0.6 + col * 6.2, 1.6 + row * 0.95
            badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.55), Inches(0.55))
            badge.fill.solid()
            badge.fill.fore_color.rgb = g.hex_to_rgb(bd.ACCENT)
            badge.line.fill.background()
            tf = badge.text_frame
            tf.paragraphs[0].text = str(i + 1)
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(16), True
            tf.paragraphs[0].font.color.rgb = g.hex_to_rgb("#FFFFFF")
            self.add_textbox(s, name, x + 0.75, y + 0.02, 5.2, 0.4, font_size=16, bold=True, color=bd.DARK)
            self.add_textbox(s, kind, x + 0.75, y + 0.4, 5.2, 0.35, font_size=11, color=bd.MUTED)

    def add_module_agenda_slide(self, d):
        """All modules at a glance: one column per day, a row per module in session order."""
        s = self._create_slide("module-agenda", d)
        self._title(s, d["title"], d.get("subtitle"))
        for col, (day, rows) in enumerate(d["days"]):
            x = 0.6 + col * 6.3
            self.add_textbox(s, day, x, 1.45, 5.9, 0.4, font_size=16, bold=True, color=bd.ACCENT)
            rule = s.shapes.add_connector(1, Inches(x), Inches(1.88), Inches(x + 5.9), Inches(1.88))
            rule.line.color.rgb = g.hex_to_rgb("#C9D2D6")
            for i, (num, time, name, tag) in enumerate(rows):
                y = 1.98 + i * 0.5
                self.add_textbox(s, time, x, y + 0.04, 1.2, 0.3, font_size=10, color=bd.MUTED)
                badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 1.2), Inches(y + 0.04), Inches(0.36),
                                           Inches(0.36))
                badge.fill.solid()
                badge.fill.fore_color.rgb = g.hex_to_rgb(bd.ACCENT)
                badge.line.fill.background()
                tf = badge.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.paragraphs[0].text = str(num)
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
                tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(11), True
                tf.paragraphs[0].font.color.rgb = g.hex_to_rgb("#FFFFFF")
                self.add_textbox(s, name, x + 1.68, y - 0.04, 4.25, 0.3, font_size=12, bold=True, color=bd.DARK)
                self.add_textbox(s, tag, x + 1.68, y + 0.2, 4.25, 0.28, font_size=10, color=bd.MUTED)

    def _box(self, s, x, y, w, h, fill="#EEF2F3"):
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = g.hex_to_rgb(fill)
        box.line.color.rgb = g.hex_to_rgb("#C9D2D6")
        box.adjustments[0] = 0.06
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.15)
        tf.margin_top = tf.margin_bottom = Inches(0.08)
        return tf

    def _rblock(self, s, label, text, x, y, w, size, label_color=None):
        """Label + text with `inline code`; returns the height used."""
        h = 0.3 + 0.2 * (size / 12) * (1 + len(text) // int((w - 0.2) * 140 / size))
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = label.upper()
        p.font.size, p.font.bold, p.font.name = Pt(10), True, g.FONT_FAMILY
        p.font.color.rgb = g.hex_to_rgb(label_color or bd.ACCENT)
        q = tf.add_paragraph()
        q.space_before = Pt(3)
        rich(q, text, size)
        return h

    def _code(self, s, code, source, x, y, w, max_h, size=10):
        """Highlighted code box, line breaks kept; if it doesn't fit: '…' and the file that holds it all."""
        per_line = int((w - 0.3) * 120 / size * 0.85)  # monospaced characters per line, with word-wrap slack
        line_h = size * 1.2 / 72
        budget = max(3, int((max_h - 0.2) / line_h))
        lines, rows = [], 0
        all_lines = code.splitlines()
        for i, line in enumerate(all_lines):
            need = 1 + len(line) // per_line
            if rows + need > budget - (1 if i < len(all_lines) - 1 else 0):
                break
            lines.append(line)
            rows += need
        cut = len(lines) < len(all_lines)
        h = 0.2 + line_h * (rows + cut)
        tf = self._box(s, x, y, w, h, fill="#F6F8FA")
        for i, line in enumerate(lines + (["…"] if cut else [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            highlight(p, line or " ", size) if line != "…" else rich(p, "…", size, bd.MUTED)
        if cut:
            note = s.shapes.add_textbox(Inches(x), Inches(y + h), Inches(w), Inches(0.3)).text_frame
            note.word_wrap = True
            rich(note.paragraphs[0], f"{len(all_lines) - len(lines)} more lines · full code: `{source}`", 9, bd.MUTED)
            h += 0.3
        return h

    def add_module_step_slide(self, d):
        s = self._create_slide("module-step", d)
        self._title(s, d["title"], d.get("subtitle"))
        img = d.get("image")
        w = 4.75 if img else 12.1
        y = 1.45
        if d.get("prompt"):  # the exact text to type: shown in full, never clipped
            text = d["prompt"]
            size = 13 if len(text) < 180 else 11 if len(text) < 330 else 10 if len(text) < 600 else 9
            rows = -(-len(text) // int((w - 0.3) * 120 / size * 0.85))  # monospaced, with word-wrap slack
            box_h = 0.25 + rows * size * 1.2 / 72
            self.add_textbox(s, "YOU TYPE (Genie Code / Genie)", 0.6, y, w, 0.3, font_size=10, bold=True,
                             color=bd.ACCENT)
            tf = self._box(s, 0.6, y + 0.35, w, box_h)
            p = tf.paragraphs[0]
            p.text, p.alignment = text, PP_ALIGN.LEFT
            p.font.size, p.font.name = Pt(size), MONO
            p.font.color.rgb = g.hex_to_rgb(bd.DARK)
            y += 0.5 + box_h
        if d.get("do"):
            label = "TALK TRACK" if d.get("talk") else "YOU DO"
            text = d["do"]
            size = 13 if len(text) < 180 else 11 if len(text) < 330 else 10
            box_h = 0.25 + len(text) / ((w - 0.3) * 140 / size) * size * 1.3 / 72
            self.add_textbox(s, label, 0.6, y, w, 0.3, font_size=10, bold=True, color=bd.ACCENT)
            rich(self._box(s, 0.6, y + 0.35, w, box_h).paragraphs[0], text, size)
            y += 0.5 + box_h
        if d.get("sql"):
            self.add_textbox(s, "YOU RUN (SQL editor)", 0.6, y, w, 0.3, font_size=10, bold=True, color=bd.ACCENT)
            room = 6.9 - (y + 0.35) - (1.0 if d.get("expected") else 0.2)
            y += 0.5 + self._code(s, d["sql"], d["source"], 0.6, y + 0.35, w, room)
        if d.get("expected"):
            y += 0.1 + self._rblock(s, "Expected result", d["expected"], 0.6, y, w, 12 if y < 5.2 else 10)
        if d.get("stuck"):
            self._rblock(s, "If you're stuck", d["stuck"], 0.6, min(y, 6.3), w, 10, label_color=bd.MUTED)
        if img:
            self._image(s, find_image(img), 5.55, 1.45, 7.25, 5.6)
        self.add_textbox(s, d.get("footer", ""), 0.6, 7.0, 12.0, 0.3, font_size=9, color=bd.MUTED)

    def generate(self, content):
        extra = {"module-step": self.add_module_step_slide, "module-overview": self.add_module_overview_slide,
                 "module-steps": self.add_module_steps_slide, "module-agenda": self.add_module_agenda_slide}
        for sd in content["slides"]:
            if sd.get("type") in extra:
                extra[sd["type"]](sd)
            else:
                bd.Deck.generate(self, {"slides": [sd]})


def header_field(header, name):
    for line in header:
        for part in line.split(" · "):
            m = re.match(r"\*\*(.+?):\*\*\s*(.*)", part.strip())
            if m and m.group(1).strip() == name:
                return clean(m.group(2))
    return ""


def build():
    paths = {parse(p)[0]: p for p in (ROOT / "code" / "prompts").glob("d[0-9][0-9]-*.md")}
    guides = sorted((parse(p) for p in paths.values()), key=lambda x: x[0])
    days = {}
    for num, title, header, *_ in guides:
        m = re.match(r"(Day \d+), ([\d:]+)[–-]([\d:]+)", header_field(header, "Session"))
        days.setdefault(m.group(1), []).append((tuple(int(v) for v in m.group(2).split(":")), num,
                                               f"{m.group(2)}–{m.group(3)}", title))
    agenda = [(day, [(num, time, re.sub(r"^Demo \d+:\s*", "", title).split(" — ")[0], TAGLINE[num])
                     for _, num, time, title in sorted(rows, key=lambda r: (r[0][0] + 12 * (r[0][0] < 8), r[0][1]))])
              for day, rows in sorted(days.items())]
    slides = [
        {"type": "title", "title": "Build it yourself: module by module",
         "subtitle": "Revantage workshop · every step, the prompt, the expected result", "date": "October 2026"},
        {"type": "module-agenda", "title": f"Workshop at a glance · {len(guides)} modules",
         "subtitle": "Every module ends with a working asset you built, with a reference to fall back on",
         "days": agenda},
        {"type": "timeline", "title": "How every module works", "steps": [
            {"title": "What we'll build", "description": "The finished asset you'll have at the end"},
            {"title": "How we'll get there", "description": "The steps, in order"},
            {"title": "Each step", "description": "Type the prompt, compare with the expected result"},
            {"title": "If stuck", "description": "Open the reference (ref_… / \"(Reference)\")"}]},
        {"type": "content", "title": "Placeholders in the prompts", "bullets": [
            "<catalog>.<schema>: the shared workshop data (your facilitator gives you the names)",
            "<my_schema>: your own sandbox schema; everything you build goes here",
            ":workspace_id: the number after ?o= in your browser URL",
            "End every prompt that builds something with \"run it and check the result\""]},
    ]
    unmatched = []
    for num, title, header, stages, _ in guides:
        name = re.sub(r"^Demo \d+:\s*", "", title)
        session = header_field(header, "Session")
        slides.append({"type": "section", "title": f"Module {num}: {name.split(' — ')[0]}",
                       "subtitle": name.split(" — ")[1].strip('"“”') if " — " in name else session})
        real = [st for st in stages if not re.match(r"(Reset|Stage \d+: Clean up)", st["title"])]
        plan = []
        for st in real:
            key = stage_key(st["title"])
            for i, step in enumerate(steps_of(st, paths[num]), 1):
                pidx = i if step["prompt"] else 0
                mkey = next((k for k in ((num, key, pidx), (num, key, 0), (num, key, 1)) if k in STEPS), None)
                if mkey is None:
                    unmatched.append((num, key, pidx))
                plan.append((st, step, STEPS.get(mkey), mkey))
        blocks = [("You build", header_field(header, "You build")), ("Session", session),
                  ("Where", header_field(header, "Where")), ("Reference (answer key)", header_field(header, "Reference"))]
        slides.append({"type": "module-overview", "title": f"What we'll build · {name.split(' — ')[0]}",
                       "blocks": [(l, clip(t, 230)) for l, t in blocks if t], "image": OVERVIEW[num]})
        items, seen = [], set()
        for st, step, media, _ in plan:
            if st["title"] in seen:
                continue
            seen.add(st["title"])
            kind = ("talk" if media and media[0] == "talk" else
                    "prompt" if step["prompt"] else "do")
            items.append((short(st["title"]), {"prompt": "Genie Code / Genie prompt", "do": "Click-path step",
                                               "talk": "Talk track"}[kind]))
        slides.append({"type": "module-steps", "title": f"How we'll get there · {len(items)} steps", "items": items[:10]})
        step_no = {}
        for st, step, media, key in plan:
            idx = list(dict.fromkeys(s["title"] for s, *_ in plan)).index(st["title"]) + 1
            step_no[idx] = step_no.get(idx, 0) + 1
            sub = f"{step_no[idx]}" if sum(1 for s, *_ in plan if s["title"] == st["title"]) > 1 else ""
            img = None
            if media and media[0] == "file":
                img = media[1]
            elif media and media[0] != "talk":
                img = fname(key)
            slides.append({"type": "module-step",
                           "title": f"Step {idx}{('.' + sub) if sub else ''} · {short(st['title'])}",
                           "subtitle": f"Module {num}: {name.split(' — ')[0]}",
                           "prompt": step["prompt"],
                           "do": None if step["prompt"] else step["do"],
                           "sql": step.get("sql") if not step["prompt"] else None, "source": step["source"],
                           "talk": bool(media and media[0] == "talk") and not step["prompt"],
                           "expected": clip(step["expected"] or ("Runs once the profiling monitor exists (created by the bundle in the Revantage workspace)." if media and media[0] == "talk" and step["prompt"] else ""), 380), "stuck": clip(step["stuck"], 200),
                           "image": img})
    slides.append({"type": "closing", "title": "Thank you"})
    if unmatched:
        print("WARNING: steps without a manifest entry:", unmatched)
    out = ROOT / "assets" / f"{bd.NAME}-modules.pptx"
    d = ModuleDeck(g.TEMPLATE_PATH)
    d.generate({"slides": slides})
    d.save(str(out))
    print("wrote", out, len(slides), "slides")


if __name__ == "__main__":
    build()
